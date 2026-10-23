"""Normalize retained official ECSA 2022 ballots and results, entirely offline.

Read original workbooks only. Retain every event (including failed/unopposed),
ballot name and typed first-preference meaning. Political party membership in the
master sheet is not treated as ballot endorsement. Current rosters stay separate.
"""
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import re
import unicodedata
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
RAW=ROOT/'content/places/raw'
CHECKED='2026-10-17'
INDEX='https://ecsa.sa.gov.au/?catid=9&id=1011%3A2022-council-elections-results-and-statistics&view=article'


def key(value):
    value=unicodedata.normalize('NFKD',str(value)).encode('ascii','ignore').decode().lower().replace('&',' and ')
    return ' '.join(sorted(w for w in re.findall('[a-z0-9]+',value) if w not in {'the','and','of','city','district','council','sa'}))


def contest_key(value):
    value=re.sub(r'\belection\b','',str(value),flags=re.I)
    value=re.sub(r'\bcouncillors\b','councillor',value,flags=re.I)
    value=re.sub(r'\bmayoral\b','mayor',value,flags=re.I)
    return key(value)


def slug(value):return re.sub('[^a-z0-9]+','-',unicodedata.normalize('NFKD',str(value)).encode('ascii','ignore').decode().lower()).strip('-')
def person_key(value):return re.sub('[^a-z0-9]','',unicodedata.normalize('NFKD',str(value)).encode('ascii','ignore').decode().lower())


def main():
    acquisition=json.loads((ROOT/'content/places/sources/council-expansion-acquisition.json').read_text(encoding='utf-8'))
    requests={Path(r['rawFile']).name:r for r in acquisition['requests'] if r.get('status')=='retained'}
    canonical=json.loads((ROOT/'content/places/councils.json').read_text(encoding='utf-8'))
    parents={}
    for place in canonical['places']:
        if place['jurisdiction']=='sa' and place['kind']=='council':
            for name in [place['name']]+place.get('aliases',[]):parents[key(name)]=place
    # ABS abbreviated the former District Council of Lower Eyre Peninsula as
    # Lower Eyre. The official event name is preserved on the historical event.
    parents[key('Lower Eyre Peninsula')]=parents[key('Lower Eyre')]
    sources=[]
    calendar_path=ROOT/'content/places/sources/sa-2022-calendar-web-review.json'
    calendar=json.loads(calendar_path.read_text(encoding='utf-8'))
    calendar_sid='council-sa-2022-postal-close-review'
    sources.append({'id':calendar_sid,'title':'ECSA 2022 postal voting close and separate count start',
                    'url':calendar['url'],'checked':CHECKED,'rawFile':calendar_path.relative_to(ROOT).as_posix(),
                    'sha256':hashlib.sha256(calendar_path.read_bytes()).hexdigest(),
                    'licence':'Minimal election-calendar facts; publisher rights retained.'})

    def source(path,label):
        request=requests[path.name]
        digest=hashlib.sha256(path.read_bytes()).hexdigest()
        if digest!=request['sha256']:raise ValueError('Retained source hash changed: '+str(path))
        sid='council-sa-2022-'+slug(label)
        if not any(s['id']==sid for s in sources):
            sources.append({'id':sid,'title':'ECSA 2022 council election '+label,'url':request['url'],'checked':CHECKED,
                            'rawFile':path.relative_to(ROOT).as_posix(),'sha256':digest,
                            'licence':'Minimal event-specific public candidate, result and typed vote facts. Full copyrighted workbook remains local-only.'})
        return sid,request['url']

    control_path=RAW/'council-expansion-sa-2022-vacancies-candidates-and-election-status-by-council.xlsx'
    control_sid,control_url=source(control_path,'contest-controls')
    controls=list(openpyxl.load_workbook(control_path,read_only=True,data_only=True).active.values)[10:]
    if len(controls)!=222:raise ValueError('Expected222 official contest controls')
    events={};patches={};audit=[]
    for council,title,seats,count,status,supplementary,comment in controls:
        parent=parents[key(council)]
        pair=(parent['id'],contest_key(title))
        if pair in events:raise ValueError('Duplicate control identity: '+str(pair))
        event={'id':parent['id']+'-ordinary-2022-'+slug(title),'label':'2022 '+council+' — '+title,
               'kind':'ordinary','status':'historical-'+status.lower(),'date':'2022-11-10',
               'dateTime':'2022-11-10T17:00:00+10:30','milestoneLabel':'postal voting closes',
               'dateBasis':'ECSA ordinary2022 postal voting closed 5pm10 November 2022; counting began 12November. Unopposed/failed offices have no invented poll or votes. Later recounts, court corrections and current office are separate.',
               'timezone':'Australia/Adelaide','sourceUrl':control_url,'resultsUrl':INDEX,
               'boundaryEdition':'2022 local election arrangement; delivered current geometry is a separately dated snapshot and is not asserted to represent2022 boundaries.',
               'contest':title,'electionCouncilName':council,'vacancies':seats,'expectedCandidateCount':count,
               'supplementaryRequired':supplementary=='Yes','candidates':[]}
        if comment:event['sourceNote']=comment
        events[pair]=event
        patch=patches.setdefault(parent['id'],{'id':parent['id'],'elections':[],'sources':[control_sid,calendar_sid],'coverageNotes':[]})
        patch['elections'].append(event)
    master_path=RAW/'council-expansion-sa-2022-details-of-candidates-and-elected-members-by-council.xlsx'
    master_sid,master_url=source(master_path,'candidate-and-historical-elected-master')
    master=list(openpyxl.load_workbook(master_path,read_only=True,data_only=True).active.values)[8:]
    if len(master)!=1256:raise ValueError('Expected1256 official candidate records')
    for row in master:
        surname,given,ballot_surname,ballot_given,council,position,ward,order,recontest,membership,elected=row
        parent=parents[key(council)]
        title=ward if position=='Ward Councillor' else ('Lord Mayor' if council=='Adelaide' and position=='Mayor' else position)
        event=events[(parent['id'],contest_key(title))]
        name=ballot_surname+', '+ballot_given
        candidate={'name':name,'status':str(elected).lower()+' in this2022 event','contest':event['contest'],
                   'sourceId':master_sid,'sourceUrl':master_url,'profileUrl':INDEX}
        if isinstance(order,(int,float)):candidate['ballotOrder']=int(order)
        if membership:
            candidate['partyMembershipAtElection']=membership
            candidate['affiliationNote']='The ECSA master records party membership; this is not an assertion of party endorsement or a ballot affiliation.'
        event['candidates'].append(candidate)
        patches[parent['id']]['sources'].append(master_sid)
    for pair,event in events.items():
        if len(event['candidates'])!=event['expectedCandidateCount']:
            raise ValueError('Master/control candidate mismatch: '+str(pair)+' '+str((len(event['candidates']),event['expectedCandidateCount'])))
    mapped=set();issues=[]
    for path in sorted(RAW.glob('council-expansion-sa-2022-*-results-statistics-*.xlsx')):
        workbook=openpyxl.load_workbook(path,read_only=True,data_only=True)
        general=next(iter(workbook));council=list(general.values)[1][0]
        parent=parents[key(council)]
        sid,url=source(path,parent['id']+'-results')
        patches[parent['id']]['sources'].append(sid)
        for sheet in list(workbook)[1:]:
            rows=list(sheet.values)
            heading=rows[1][0].split(',',1)[-1].strip()
            printed_heading=heading
            if parent['id']=='sa-council-40700' and sheet.title=='Eastwood & Glenunga ward' and heading=='Eastside & Glenunga ward':
                # Only this exact workbook typo is reconciled. The sheet tab,
                # master and independent contest-control workbook all agree.
                heading=sheet.title
            pair=(parent['id'],contest_key(heading));event=events.get(pair)
            if not event:
                issues.append({'kind':'unjoined-result-sheet','parent':parent['id'],'sheet':sheet.title,'heading':heading});continue
            if pair in mapped:raise ValueError('Result contest mapped twice: '+str(pair))
            mapped.add(pair)
            if printed_heading!=heading:event['sourceNotes']=['The Burnside result worksheet heading prints Eastside & Glenunga; its sheet tab, candidate master and separate contest-control workbook identify Eastwood & Glenunga. This exact supported discrepancy is retained and reconciled.']
            pos=next((i for i,r in enumerate(rows) if any('first preference votes (no.)'==str(v).strip().lower() for v in r)),None)
            if pos is None:
                issues.append({'kind':'no-first-preference-table','parent':parent['id'],'sheet':sheet.title});continue
            summary=[]
            for row in rows[pos+1:]:
                if not row[0] or not isinstance(row[0],str) or ',' not in row[0] or not isinstance(row[1],(int,float)):
                    break
                summary.append(row)
            if len(summary)!=event['expectedCandidateCount']:
                issues.append({'kind':'result-candidate-count','pair':pair,'actual':len(summary),'expected':event['expectedCandidateCount']});continue
            by_name={person_key(c['name']):c for c in event['candidates']}
            for row in summary:
                candidate=by_name.get(person_key(row[0]))
                if not candidate:
                    issues.append({'kind':'unjoined-candidate-name','pair':pair,'name':row[0]});continue
                candidate.update(votes=row[1],voteMeasure='first preference votes',sourceId=sid,sourceUrl=url)
                if str(row[4]).strip().lower() in ['elected','not elected']:
                    candidate['masterHistoricalStatus']=candidate['status']
                    candidate['status']=str(row[4]).strip().lower()+' in this2022 result workbook'
            formal=next((r[1] for r in rows if str(r[0]).strip()=='Formal votes'),None)
            actual=sum(row[1] for row in summary)
            if formal is not None and actual!=formal:issues.append({'kind':'first-preference-total','pair':pair,'actual':actual,'expected':formal})
            event.update(sourceUrl=url,resultsUrl=url)
            audit.append({'place':parent['id'],'event':event['id'],'sheet':sheet.title,'candidateCount':len(summary),'firstPreferences':actual,'formalVotes':formal,'sourceId':sid})
    missing=[pair for pair,event in events.items() if event['status']=='historical-contested' and pair not in mapped]
    if missing or issues:
        print(json.dumps({'missing':missing,'issues':issues},ensure_ascii=True,indent=2));raise ValueError('Source reconciliation incomplete; no supplement written')
    for patch in list(patches.values()):
        patch['sources']=list(dict.fromkeys(patch['sources']))
        patch['coverageNotes']=['All 2022 ordinary-event candidates come from ECSA master/control workbooks; contested first preferences come from the separate official council results workbook. Source-published historical outcome corrections are retained as published. These records do not establish current office or completion of later by-election/recount histories.','Party membership in the historical master is retained separately and is not treated as party endorsement.']
    cancelled_parent=parents[key('Coober Pedy')]
    patches[cancelled_parent['id']]={'id':cancelled_parent['id'],'sources':[control_sid,master_sid],
        'elections':[{'id':cancelled_parent['id']+'-ordinary-2022-cancelled','label':'2022 Coober Pedy ordinary election cancelled',
                      'kind':'ordinary','status':'cancelled','date':None,'timezone':'Australia/Adelaide',
                      'dateBasis':'Both ECSA 2022 official master and contest-control workbooks explicitly say the Coober Pedy periodic elections were cancelled. No contested ballot, candidate or result is invented.',
                      'sourceUrl':control_url,'resultsUrl':INDEX,'candidates':[]}],
        'coverageNotes':['The2022 periodic election was cancelled, as explicitly stated in ECSA master/control workbooks. This historical cancellation does not establish the current governing arrangement.']}
    # Ward projection follows only exact parent + historical ward label joins.
    # The event retains its2022 edition caveat, rather than claiming current
    # polygons are necessarily the polygons under which this contest occurred.
    ward_events=0
    for ward in canonical['places']:
        if ward['jurisdiction']!='sa' or ward['kind']!='ward':continue
        event=events.get((ward['parentId'],contest_key(ward['name'])))
        if not event:continue
        projected={**event,'id':ward['id']+'-ordinary-2022'}
        patches[ward['id']]={'id':ward['id'],'elections':[projected],'sources':patches[ward['parentId']]['sources'],
                             'coverageNotes':['The2022 ward contest has an exact parent-and-name correspondence. Its historical arrangement is kept separate from the current delivered ward geometry; no boundary equivalence is asserted.']}
        ward_events+=1
    result={'schemaVersion':1,'checked':CHECKED,'sources':sources,'places':list(patches.values()),'checks':audit,
            'coverage':{'councils':len({p for p,c in events}),'parentContests':len(events),'parentContestants':len(master),
                        'contested':sum(e['status']=='historical-contested' for e in events.values()),
                        'uncontested':sum(e['status']=='historical-uncontested' for e in events.values()),
                        'failed':sum(e['status']=='historical-failed' for e in events.values()),'projectedWardContests':ward_events}}
    (ROOT/'content/places/sa-council-ballots.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result['coverage']))


if __name__=='__main__':main()
